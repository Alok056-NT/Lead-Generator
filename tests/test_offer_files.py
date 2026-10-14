"""Offer file sources: TXT, MD, DOCX and PDF (fixtures are generated here, never committed)."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from pypdf import PdfWriter

from leadgen.services import offer_files
from leadgen.services.offer_files import FileSourceError, read_file_text


def make_pdf(lines: list[str]) -> bytes:
    """A one-page PDF with real text, written by hand (pypdf cannot write text)."""
    body = "BT /F1 12 Tf 72 720 Td " + " ".join(f"({line}) Tj 0 -16 Td" for line in lines) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(body)} >>\nstream\n{body}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{obj}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    trailer = f"<< /Size {len(objects) + 1} /Root 1 0 R >>"
    out += f"trailer\n{trailer}\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


# --- plain text -------------------------------------------------------------------------------


def test_txt_and_md(tmp_path):
    (tmp_path / "a.txt").write_text("Line one\r\n\r\n\r\n\r\nLine two   \n", encoding="utf-8")
    (tmp_path / "b.MD").write_text("# Title\n\ntext", encoding="utf-8")
    assert read_file_text(tmp_path / "a.txt") == "Line one\n\nLine two"
    assert read_file_text(tmp_path / "b.MD") == "# Title\n\ntext"


def test_text_encodings(tmp_path):
    (tmp_path / "bom.txt").write_bytes(b"\xef\xbb\xbfCaf\xc3\xa9")
    (tmp_path / "old.txt").write_bytes("Café".encode("cp1252"))
    assert read_file_text(tmp_path / "bom.txt") == "Café"
    assert read_file_text(tmp_path / "old.txt") == "Café"


def test_rejections(tmp_path):
    def code_for(path: Path) -> str:
        with pytest.raises(FileSourceError) as info:
            read_file_text(path)
        return info.value.code

    assert code_for(tmp_path / "missing.txt") == "file_not_found"
    assert code_for(tmp_path) == "file_not_found"  # a folder is not a file
    (tmp_path / "x.xlsx").write_bytes(b"x")
    assert code_for(tmp_path / "x.xlsx") == "unsupported_type"
    (tmp_path / "x.doc").write_bytes(b"x")
    assert code_for(tmp_path / "x.doc") == "unsupported_type"
    (tmp_path / "noext").write_bytes(b"x")
    assert code_for(tmp_path / "noext") == "unsupported_type"
    (tmp_path / "empty.txt").write_text(" \n ", encoding="utf-8")
    assert code_for(tmp_path / "empty.txt") == "no_text"


def test_size_limit(tmp_path, monkeypatch):
    monkeypatch.setattr(offer_files, "MAX_FILE_BYTES", 10)
    (tmp_path / "big.txt").write_text("x" * 11, encoding="utf-8")
    with pytest.raises(FileSourceError) as info:
        read_file_text(tmp_path / "big.txt")
    assert info.value.code == "file_too_large"
    (tmp_path / "ok.txt").write_text("x" * 10, encoding="utf-8")
    assert read_file_text(tmp_path / "ok.txt") == "x" * 10


# --- DOCX -------------------------------------------------------------------------------------


def build_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("Northwind Training", level=1)
    doc.add_paragraph("We run safety courses for warehouses.")
    doc.add_paragraph("Forklift certification", style="List Bullet")
    para = doc.add_paragraph("Our success rate is ")
    run = para.add_run("[CONFIRM: 98%]")
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    para.add_run(" across clients.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Course"
    table.cell(0, 1).text = "Length"
    table.cell(1, 0).text = "First aid"
    table.cell(1, 1).text = "1 day"
    merged = doc.add_table(rows=1, cols=3)
    merged.cell(0, 0).merge(merged.cell(0, 1)).text = "Wide header"
    merged.cell(0, 2).text = "Side"
    doc.add_paragraph("Closing line")
    doc.save(str(path))


def test_docx_text_tables_and_highlights(tmp_path):
    path = tmp_path / "offer.docx"
    build_docx(path)
    lines = read_file_text(path).splitlines()
    assert lines[0] == "## Northwind Training"
    assert "We run safety courses for warehouses." in lines
    assert "- Forklift certification" in lines
    assert "Our success rate is [HIGHLIGHT: [CONFIRM: 98%]] across clients." in lines
    assert "Course | Length" in lines and "First aid | 1 day" in lines
    assert "Wide header | Side" in lines  # merged cell written once
    assert lines.index("Course | Length") < lines.index("Closing line")  # document order


def test_docx_highlight_none_is_not_highlighted(tmp_path):
    """`<w:highlight w:val="none"/>` is valid Word output; it must not make the file unreadable."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    path = tmp_path / "none.docx"
    doc = Document()
    run = doc.add_paragraph().add_run("Plain text with no highlight.")
    highlight = OxmlElement("w:highlight")
    highlight.set(qn("w:val"), "none")
    run._r.get_or_add_rPr().append(highlight)
    doc.save(str(path))
    text = read_file_text(path)
    assert text == "Plain text with no highlight."
    assert "[HIGHLIGHT" not in text


def test_docx_corrupt(tmp_path):
    (tmp_path / "bad.docx").write_bytes(b"not a zip")
    with pytest.raises(FileSourceError) as info:
        read_file_text(tmp_path / "bad.docx")
    assert info.value.code == "unreadable"


def test_docx_without_text(tmp_path):
    Document().save(str(tmp_path / "blank.docx"))
    with pytest.raises(FileSourceError) as info:
        read_file_text(tmp_path / "blank.docx")
    assert info.value.code == "no_text"


# --- PDF --------------------------------------------------------------------------------------


def test_pdf_text(tmp_path):
    (tmp_path / "offer.pdf").write_bytes(make_pdf(["Northwind Training", "Safety courses"]))
    text = read_file_text(tmp_path / "offer.pdf")
    assert "Northwind Training" in text and "Safety courses" in text


def test_pdf_without_text_layer(tmp_path):
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with (tmp_path / "scan.pdf").open("wb") as handle:
        writer.write(handle)
    with pytest.raises(FileSourceError) as info:
        read_file_text(tmp_path / "scan.pdf")
    assert info.value.code == "no_text"


def test_pdf_encrypted_and_corrupt(tmp_path):
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("user-pass", "owner-pass")
    with (tmp_path / "locked.pdf").open("wb") as handle:
        writer.write(handle)
    with pytest.raises(FileSourceError) as info:
        read_file_text(tmp_path / "locked.pdf")
    assert info.value.code == "password_protected"

    (tmp_path / "bad.pdf").write_bytes(b"this is not a pdf at all")
    with pytest.raises(FileSourceError) as info:
        read_file_text(tmp_path / "bad.pdf")
    assert info.value.code == "unreadable"
