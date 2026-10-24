"""Importing old contacts: files, cleaning, modes, dry runs, the v3 migration (T24).

Every person, company and address in this file is invented (example.com / .example domains).
No network is used. Files are written into pytest's temporary folder.
"""

from __future__ import annotations

import csv
import shutil
import sqlite3
from datetime import UTC, date, datetime, timedelta
from pathlib import Path

import openpyxl
import pytest
from qa_world import COMPANY, add_company, add_person, make_campaign

from leadgen import db
from leadgen.services import (
    browse,
    companies,
    contacts_import,
    drafting,
    people,
    projects,
    qa_leads,
    suppression,
)
from leadgen.services.campaign import get_campaign
from leadgen.services.contacts_import import (
    COLUMNS,
    ContactsImportError,
    import_contacts,
    read_contacts,
)

TODAY = datetime.now(UTC).date()


def days_ago(n: int) -> date:
    return TODAY - timedelta(days=n)


def row(**values):
    return {c: values.get(c, "") for c in COLUMNS}


def write_csv(path, rows, header=COLUMNS, delimiter=",", encoding="utf-8-sig"):
    with open(path, "w", newline="", encoding=encoding) as handle:
        writer = csv.writer(handle, delimiter=delimiter)
        writer.writerow(header)
        for r in rows:
            writer.writerow([r.get(c, "") for c in header] if isinstance(r, dict) else r)
    return path


def write_xlsx(path, rows, header=COLUMNS, sheet="Contacts"):
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = sheet
    ws.append(list(header))
    for r in rows:
        ws.append(r if isinstance(r, (list, tuple)) else [r.get(c, "") for c in header])
    book.save(path)
    return path


def table_names(conn):
    return {r[0] for r in conn.execute("SELECT name FROM sqlite_master")}


def count(table, where="1=1", params=()):
    with db.session() as conn:
        return conn.execute(f"SELECT COUNT(*) FROM {table} WHERE {where}", params).fetchone()[0]


def table_counts():
    return {t: count(t) for t in ("contact_log", "known_contacts", "suppression")}


@pytest.fixture
def project(ready_db):
    return projects.create_project("Northwind Training")


@pytest.fixture
def pid(project):
    return project["id"]


def by_name(preview):
    return {c.name: c for c in preview.contacts}


# --- template ---------------------------------------------------------------------------------


def test_the_shipped_csv_template_is_the_header_only(tmp_path):
    data = (contacts_import._csv_template_bytes()).decode("utf-8-sig")
    assert data.strip() == ",".join(COLUMNS)
    written = contacts_import.export_template(tmp_path / "t.csv")
    assert written.read_bytes() == contacts_import._csv_template_bytes()
    assert contacts_import.export_template(tmp_path, "csv").name == "contacts_template.csv"
    with pytest.raises(ValueError):
        contacts_import.export_template(tmp_path / "x.pdf", "pdf")


def test_the_excel_template_has_both_sheets_and_one_guide_line_per_column(tmp_path):
    path = contacts_import.export_template(tmp_path / "t.xlsx", "xlsx")
    book = openpyxl.load_workbook(path)
    assert book.sheetnames == ["Contacts", "How to fill this in"]
    assert [c.value for c in book["Contacts"][1]] == list(COLUMNS)
    guide = [r for r in book["How to fill this in"].iter_rows(values_only=True)]
    assert [g[0] for g in guide[1 : 1 + len(COLUMNS)]] == list(COLUMNS)
    assert all(g[2] for g in guide[1 : 1 + len(COLUMNS)])
    assert set(contacts_import.COLUMN_GUIDE) == set(COLUMNS)
    # an empty template imports nothing and is not an error
    empty = read_contacts(path)
    assert empty.error == "" and empty.contacts == [] and empty.rows_read == 0


def test_template_info_writes_both_files_into_the_workspace(ready_db, workspace):
    info = contacts_import.template_info()
    assert info["ok"] and [c["name"] for c in info["columns"]] == list(COLUMNS)
    for key in ("csv_path", "xlsx_path"):
        assert (workspace / "templates").resolve() in Path(info[key]).resolve().parents
    assert "do_not_contact" in info["relationship_values"]
    assert contacts_import.template_info()["ok"]  # writing again is fine


def test_template_round_trip_csv(tmp_path):
    template = contacts_import.export_template(tmp_path / "t.csv")
    with open(template, "a", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Aria", "Tan", "", "aria.tan@northwind.example", "", "Northwind Logistics",
                         "northwind.example", "Head of Operations", "Singapore", "2026-09-30",
                         "email", "customer", "Met at the expo"])  # fmt: skip
        writer.writerow(["", "", "Omar Haddad", "", "linkedin.com/in/omar-haddad",
                         "Contoso Freight", "", "", "", "30/09/2026", "linkedin", "",
                         ""])  # fmt: skip
        writer.writerow(["Hana", "Reyes", "", "", "", "Fabrikam Cargo", "", "", "", "", "", "",
                         ""])  # fmt: skip
    preview = read_contacts(template)
    assert preview.error == "" and preview.rows_read == 3 and preview.valid_rows == 3
    people_ = by_name(preview)
    aria = people_["Aria Tan"]
    assert aria.email == "aria.tan@northwind.example" and aria.relationship == "customer"
    assert aria.last_contacted == "2026-09-30" and aria.company_domain == "northwind.example"
    assert aria.notes == "Met at the expo" and aria.channel == "email"
    omar = people_["Omar Haddad"]  # full_name used because first/last are empty
    assert (omar.first_name, omar.last_name) == ("Omar", "Haddad")
    assert omar.linkedin_url == "https://linkedin.com/in/omar-haddad"  # as the finder keeps it
    assert omar.last_contacted == "2026-09-30" and omar.relationship == "contacted"
    hana = people_["Hana Reyes"]
    assert hana.last_contacted == ""  # empty: treated as contacted at import time
    assert [w["column"] for w in preview.warnings] == ["row"]  # hana: name + company only


def test_template_round_trip_xlsx(tmp_path):
    path = contacts_import.export_template(tmp_path / "t.xlsx", "xlsx")
    book = openpyxl.load_workbook(path)
    ws = book["Contacts"]
    ws.append(["Aria", "Tan", "", "aria.tan@northwind.example", "", "Northwind Logistics"])
    ws.append(["Omar", "Haddad", "", "omar@contoso.example", "", "Contoso Freight"])
    ws.append(["Hana", "Reyes", "", "hana@fabrikam.example", "", "Fabrikam Cargo"])
    book.save(path)
    preview = read_contacts(path)
    assert preview.error == "" and preview.format == "xlsx"
    assert sorted(by_name(preview)) == ["Aria Tan", "Hana Reyes", "Omar Haddad"]


# --- file formats -----------------------------------------------------------------------------

LINKEDIN_EXPORT = (
    "Notes:\n"
    '"When exporting your connection data, you may notice that some of the email addresses are '
    "missing, because only connections who allow it share them. You can learn more here, in the "
    'help centre."\n'
    "\n"
    "First Name,Last Name,URL,Email Address,Company,Position,Connected On\n"
    "Aria,Tan,https://www.linkedin.com/in/aria-tan,,Northwind Logistics,Head of Operations,"
    "14 Oct 2026\n"
    "Omar,Haddad,https://sg.linkedin.com/in/omar-haddad/,omar@contoso.example,Contoso Freight,"
    "Director,03 Jan 2021\n"
    ",,https://www.linkedin.com/in/ghost,,,,\n"
)


def test_linkedin_connections_csv_with_the_notes_header(tmp_path):
    path = tmp_path / "Connections.csv"
    path.write_text(LINKEDIN_EXPORT, encoding="utf-8")
    preview = read_contacts(path)
    assert preview.error == "" and preview.source == "linkedin_connections"
    assert preview.rows_read == 3 and preview.valid_rows == 3
    contacts = by_name(preview)
    aria = contacts["Aria Tan"]
    assert aria.linkedin_url == "https://www.linkedin.com/in/aria-tan"
    assert aria.company == "Northwind Logistics" and aria.title == "Head of Operations"
    assert aria.last_contacted == ""  # being connected is not the same as having been contacted
    assert aria.notes == "LinkedIn connection since 2026-10-14."
    assert contacts["Omar Haddad"].linkedin_url == "https://sg.linkedin.com/in/omar-haddad"
    assert contacts["Omar Haddad"].email == "omar@contoso.example"
    nameless = [c for c in preview.contacts if not c.name]
    assert len(nameless) == 1 and nameless[0].linkedin_url.endswith("/in/ghost")
    keys = aria.keys()
    assert keys["key_linkedin"] == "in/aria-tan"  # the contact log stores in/<slug>


def test_cp1252_file_with_accents(tmp_path):
    path = write_csv(
        tmp_path / "old.csv",
        [row(first_name="José", last_name="Müller", company="Café Zoë AG",
             email="jose@cafe-zoe.example")],  # fmt: skip
        encoding="cp1252",
    )
    assert b"Jos\xe9" in path.read_bytes()  # really not UTF-8
    preview = read_contacts(path)
    assert preview.error == "" and preview.contacts[0].first_name == "José"
    first = preview.contacts[0]
    assert first.last_name == "Müller" and first.company == "Café Zoë AG"


def test_semicolon_and_tab_separated_files(tmp_path):
    rows = [row(first_name="Aria", last_name="Tan", email="aria@northwind.example",
                company="Northwind Logistics", notes="one, two")]  # fmt: skip
    for name, delimiter in (("semi.csv", ";"), ("tab.csv", "\t"), ("comma.csv", ",")):
        preview = read_contacts(write_csv(tmp_path / name, rows, delimiter=delimiter))
        assert preview.error == "" and len(preview.contacts) == 1, name
        assert preview.contacts[0].notes == "one, two", name


def test_utf16_excel_text_export(tmp_path):
    path = tmp_path / "u16.csv"
    path.write_text("first_name\tlast_name\temail\nAria\tTan\taria@northwind.example\n",
                    encoding="utf-16")  # fmt: skip
    preview = read_contacts(path)
    assert preview.error == "" and preview.contacts[0].email == "aria@northwind.example"


def test_headers_are_case_and_space_insensitive_and_order_is_free(tmp_path):
    path = write_csv(
        tmp_path / "x.csv",
        [["aria@northwind.example", "Tan", "Aria", "Northwind Logistics", "ignore me"]],
        header=["E-Mail", " LAST NAME ", "First  Name", "Company", "Favourite colour"],
    )
    preview = read_contacts(path)
    assert preview.error == "" and preview.contacts[0].name == "Aria Tan"
    assert preview.unknown_columns == ["Favourite colour"]
    assert preview.summary()["unknown_columns"] == ["Favourite colour"]


def test_xlsx_with_excel_date_cells_and_blank_rows(tmp_path):
    path = write_xlsx(
        tmp_path / "book.xlsx",
        [
            ["Aria", "Tan", "", "aria@northwind.example", "", "Northwind Logistics", "", "", "",
             datetime(2026, 9, 30, 0, 0), "", "replied", ""],
            [None] * len(COLUMNS),
            [None] * len(COLUMNS),
            ["Omar", "Haddad", "", "omar@contoso.example", "", "Contoso Freight", "", "", "",
             date(2026, 8, 1), "", "", ""],
            ["", "", "", "", "", "", "", "", "", "", "", "", ""],
            ["Hana", "Reyes", "", "hana@fabrikam.example", "", "Fabrikam Cargo", "", "", "",
             "15/07/2026", "", "", ""],
        ],
    )  # fmt: skip
    preview = read_contacts(path)
    assert preview.error == "" and preview.rows_read == 3 and preview.valid_rows == 3
    contacts = by_name(preview)
    assert contacts["Aria Tan"].last_contacted == "2026-09-30"
    assert contacts["Aria Tan"].relationship == "replied"
    assert contacts["Omar Haddad"].last_contacted == "2026-08-01"
    assert contacts["Hana Reyes"].last_contacted == "2026-07-15"
    assert contacts["Hana Reyes"].rows == [7]  # row numbers are the ones Excel shows


def test_xlsx_uses_the_contacts_sheet_even_when_it_is_not_first(tmp_path):
    path = tmp_path / "two.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "Cover"
    book.active.append(["Read me"])
    sheet = book.create_sheet("contacts")
    sheet.append(["first_name", "last_name", "email"])
    sheet.append(["Aria", "Tan", "aria@northwind.example"])
    book.save(path)
    assert read_contacts(path).contacts[0].name == "Aria Tan"


def test_xlsx_every_row_is_read(tmp_path):
    path = write_xlsx(
        tmp_path / "x.xlsx",
        [row(first_name=f"P{i}", last_name="Tan", email=f"p{i}@a.example") for i in range(30)],
    )
    assert len(read_contacts(path).contacts) == 30


# --- bad input never raises -------------------------------------------------------------------


def test_bad_rows_are_reported_not_raised(tmp_path):
    today = TODAY.isoformat()
    future = (TODAY + timedelta(days=30)).isoformat()
    path = write_csv(
        tmp_path / "messy.csv",
        [
            row(first_name="Aria", last_name="Tan", email="not-an-email",
                linkedin_url="https://www.linkedin.com/in/aria-tan"),  # bad email, good link
            row(first_name="Omar", last_name="Haddad", email="omar@contoso.example",
                linkedin_url="https://www.linkedin.com/company/contoso"),  # company page
            row(first_name="Hana", last_name="Reyes", company="Fabrikam Cargo",
                last_contacted="yesterday-ish"),  # bad date: row skipped
            row(first_name="Ines", last_name="Costa", company="Fabrikam Cargo",
                relationship="best friend"),  # bad relationship: row skipped
            row(first_name="Lena", last_name="Park"),  # nothing to match on
            row(first_name="Nia", last_name="Okoye", email="nia@fabrikam.example",
                channel="carrier pigeon", company_domain="not a domain"),
            row(first_name="Sam", last_name="Ong", email="sam@fabrikam.example",
                last_contacted=future),
            row(first_name="Tess", last_name="Wu", email="tess@fabrikam.example",
                last_contacted="1850-01-01"),  # impossible date
        ],
    )  # fmt: skip
    preview = read_contacts(path)
    assert preview.error == ""
    assert preview.rows_read == 8
    problems = {(e["row"], e["column"]) for e in preview.errors}
    assert (2, "email") in problems  # row 2 is the first data row
    assert (3, "linkedin_url") in problems
    assert (4, "last_contacted") in problems
    assert (5, "relationship") in problems
    assert (6, "row") in problems
    assert (7, "channel") in problems and (7, "company_domain") in problems
    assert (9, "last_contacted") in problems
    kept = by_name(preview)
    assert kept["Aria Tan"].email == "" and kept["Aria Tan"].linkedin_url  # field dropped, row kept
    assert kept["Omar Haddad"].linkedin_url == "" and kept["Omar Haddad"].email
    assert "Hana Reyes" not in kept and "Ines Costa" not in kept and "Lena Park" not in kept
    assert "Tess Wu" not in kept
    assert kept["Sam Ong"].last_contacted == today  # a date in the future becomes today
    assert any(w["column"] == "last_contacted" for w in preview.warnings)
    assert preview.rows_skipped == 4 and preview.valid_rows == 4
    assert all({"row", "column", "problem"} == set(e) for e in preview.errors)


def test_a_row_with_no_email_or_link_gets_a_warning(tmp_path):
    path = write_csv(
        tmp_path / "w.csv", [row(first_name="Hana", last_name="Reyes", company="Fabrikam Cargo")]
    )
    preview = read_contacts(path)
    assert preview.valid_rows == 1 and preview.errors == []
    assert "name and company" in preview.warnings[0]["problem"]


def test_unusable_files_give_a_message_not_an_exception(tmp_path, monkeypatch):
    assert "no file" in read_contacts(tmp_path / "missing.csv").error.lower()
    pdf = tmp_path / "list.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    assert ".csv or .xlsx" in read_contacts(pdf).error
    xls = tmp_path / "old.xls"
    xls.write_bytes(b"\xd0\xcf\x11\xe0")
    assert ".csv or .xlsx" in read_contacts(xls).error
    bad = tmp_path / "broken.xlsx"
    bad.write_bytes(b"this is not a zip file" * 10)
    assert "could not be read" in read_contacts(bad).error
    empty = tmp_path / "empty.csv"
    empty.write_text("", encoding="utf-8")
    assert "No header row" in read_contacts(empty).error
    junk = tmp_path / "junk.csv"
    junk.write_text("hello\nworld\n", encoding="utf-8")
    assert "No header row" in read_contacts(junk).error
    nul = tmp_path / "nul.csv"
    nul.write_bytes(b"first_name,email\x00\nAria,aria@northwind.example\n")
    cleaned = read_contacts(nul)
    assert cleaned.error == "" and cleaned.contacts[0].email == "aria@northwind.example"
    big = write_csv(tmp_path / "big.csv", [row(first_name="A", last_name="B", email="a@b.example")])
    monkeypatch.setattr(contacts_import, "MAX_BYTES", 10)
    assert "larger than" in read_contacts(big).error
    for item in (read_contacts(tmp_path / "missing.csv"), read_contacts(big)):
        assert item.contacts == []


def test_more_than_50000_rows_are_refused(tmp_path):
    lines = ["first_name,email"] + [f"P{i},p{i}@northwind.example" for i in range(50_001)]
    path = tmp_path / "huge.csv"
    path.write_text("\n".join(lines), encoding="utf-8")
    preview = read_contacts(path)
    assert "more than 50,000 rows" in preview.error and preview.contacts == []
    exactly = tmp_path / "exact.csv"
    exactly.write_text("\n".join(lines[:-1]), encoding="utf-8")
    ok = read_contacts(exactly)
    assert ok.error == "" and ok.rows_read == 50_000 and len(ok.contacts) == 50_000


def test_the_row_limit_applies_to_excel_too(tmp_path, monkeypatch):
    monkeypatch.setattr(contacts_import, "MAX_ROWS", 3)
    rows = [row(first_name=f"P{i}", last_name="T", email=f"p{i}@a.example") for i in range(4)]
    assert "more than 3 rows" in read_contacts(write_xlsx(tmp_path / "x.xlsx", rows)).error


def test_repeated_rows_of_one_person_are_merged(tmp_path):
    path = write_csv(
        tmp_path / "dups.csv",
        [
            row(first_name="Aria", last_name="Tan", email="Aria.Tan@Northwind.example",
                company="Northwind Logistics", last_contacted="2026-03-01",
                relationship="contacted", notes="first"),
            row(first_name="Aria", last_name="Tan", company="Northwind Logistics",
                linkedin_url="https://sg.linkedin.com/in/aria-tan/",
                last_contacted="2026-09-01", relationship="customer", channel="email"),
            row(email="aria.tan@northwind.example", title="Head of Ops", channel="linkedin",
                notes="second"),
            row(first_name="Omar", last_name="Haddad", company="Contoso Freight"),
            row(first_name="omar", last_name="haddad", company="CONTOSO freight!"),
        ],
    )  # fmt: skip
    preview = read_contacts(path)
    assert preview.rows_read == 5 and preview.valid_rows == 5 and preview.duplicates == 3
    assert len(preview.contacts) == 2
    aria = by_name(preview)["Aria Tan"]
    assert aria.email == "aria.tan@northwind.example"
    assert aria.linkedin_url == "https://sg.linkedin.com/in/aria-tan"
    assert aria.keys()["key_linkedin"] == "in/aria-tan"
    assert aria.last_contacted == "2026-09-01" and aria.relationship == "customer"
    assert aria.title == "Head of Ops" and aria.channel == "both"
    assert aria.notes == "first | second" and aria.rows == [2, 3, 4]


def test_relationship_and_channel_spellings(tmp_path):
    path = write_csv(
        tmp_path / "s.csv",
        [
            row(email="a@x.example", relationship="Do Not Contact"),
            row(email="b@x.example", relationship="DNC"),
            row(email="c@x.example", relationship="Client", channel="E-mail"),
            row(email="d@x.example", relationship=" Replied ", channel="LinkedIn"),
        ],
    )
    got = {c.email[0]: (c.relationship, c.channel) for c in read_contacts(path).contacts}
    assert got == {
        "a": ("do_not_contact", ""), "b": ("do_not_contact", ""),
        "c": ("customer", "email"), "d": ("replied", "linkedin"),
    }  # fmt: skip


# --- the three modes --------------------------------------------------------------------------


def people_file(tmp_path, **extra):
    rows = [
        row(first_name="Aria", last_name="Tan", email="aria.tan@northwind.example",
            company=COMPANY, last_contacted=days_ago(10).isoformat(), relationship="replied"),
        row(first_name="Omar", last_name="Haddad", linkedin_url="https://www.linkedin.com/in/omar-haddad",
            company="Contoso Freight", last_contacted=days_ago(400).isoformat()),
        row(first_name="Hana", last_name="Reyes", company="Fabrikam Cargo",
            relationship="customer"),
        row(first_name="Ines", last_name="Costa", email="ines@fabrikam.example",
            linkedin_url="https://www.linkedin.com/in/ines-costa", relationship="do_not_contact"),
        row(first_name="Sam", last_name="Ong", company="Fabrikam Cargo",
            relationship="do_not_contact"),
    ]  # fmt: skip
    return write_csv(tmp_path / "old.csv", rows)


def test_dry_run_writes_nothing_and_reports_counts(tmp_path, pid):
    path = people_file(tmp_path)
    before = table_counts()
    result = import_contacts(pid, path, "both")  # dry_run defaults to true
    assert table_counts() == before
    assert result["dry_run"] is True and result["ok"] is True
    assert result["rows_read"] == 5 and result["valid_rows"] == 5 and result["contacts"] == 5
    assert result["duplicates_merged"] == 0 and result["error_count"] == 0
    assert result["would_exclude"] == 5 and result["would_signal"] == 3
    assert result["do_not_contact"] == 2 and result["errors"] == []
    assert "excluded" not in result and "signalled" not in result
    assert result["message"].startswith("Preview only, nothing was written.")
    assert "dry_run=false" in result["next_step"]
    assert result["outside_cooldown"] == 1  # Omar was last contacted 400 days ago
    assert result["do_not_contact_name_only"] == 1  # Sam has no email or link
    # also true after an exception halfway (a bad project) and when applying nothing
    with pytest.raises(projects.ProjectError):
        import_contacts(9999, path, "both", dry_run=False)
    assert table_counts() == before


def test_the_first_20_errors_are_shown(tmp_path, pid):
    rows = [row(first_name=f"P{i}", last_name="T", email=f"bad{i}", linkedin_url="https://x.example")
            for i in range(30)]  # fmt: skip
    result = import_contacts(pid, write_csv(tmp_path / "e.csv", rows), "exclude")
    assert result["error_count"] >= 30 and len(result["errors"]) == 20


def test_exclude_logs_contacts_and_l24_blocks_them(tmp_path, pid):
    result = import_contacts(pid, people_file(tmp_path), "exclude", dry_run=False)
    assert result["dry_run"] is False and result["excluded"] == 5 and "signalled" not in result
    assert count("known_contacts") == 0
    with db.session() as conn:
        log = db.query(conn, "SELECT * FROM contact_log ORDER BY id")
    assert {r["channel"] for r in log} == {"imported"}
    assert all(r["source"] == "import: old.csv" and r["campaign_id"] is None for r in log)
    aria = next(r for r in log if r["key_email"] == "aria.tan@northwind.example")
    assert aria["at"].startswith(days_ago(10).isoformat())
    assert aria["key_name_company"] == "aria tan|northwind logistics"
    assert next(r for r in log if r["key_linkedin"] == "in/omar-haddad")  # stored as in/<slug>
    hana = next(r for r in log if r["key_name_company"].startswith("hana reyes"))
    assert hana["at"].startswith(TODAY.isoformat())  # empty date: contacted at import time


def test_l24_blocks_an_imported_person_inside_the_cooldown(ready_db, tmp_path):
    campaign = make_campaign()
    cid, project_id = campaign["id"], get_campaign(campaign["id"])["project_id"]
    company = add_company(cid)
    blocked = add_person(cid, company, "Aria", "Tan")
    old = add_person(cid, company, "Omar", "Haddad")
    free = add_person(cid, company, "Priya", "Nair")
    path = write_csv(
        tmp_path / "old.csv",
        [
            row(first_name="Aria", last_name="Tan", company=COMPANY,
                last_contacted=days_ago(10).isoformat()),
            row(first_name="Omar", last_name="Haddad", company=COMPANY,
                last_contacted=days_ago(400).isoformat()),
        ],
    )  # fmt: skip
    import_contacts(project_id, path, "exclude", dry_run=False)

    def issues():
        data = qa_leads.load_leads(get_campaign(cid))
        return {
            p: {i.code: i for i in found} for p, found in qa_leads.check_all(data).items()
        }

    found = issues()
    assert found[blocked]["L24"].severity == "block"
    assert "imported" in found[blocked]["L24"].detail
    assert "L24" not in found[old]  # outside the cooldown
    assert "L24" not in found[free]


def test_do_not_contact_goes_on_the_global_list(tmp_path, pid):
    import_contacts(pid, people_file(tmp_path), "exclude", dry_run=False)
    listed = {(r["kind"], r["value"]) for r in suppression.list()}
    assert listed == {
        ("email", "ines@fabrikam.example"),
        ("linkedin", "https://www.linkedin.com/in/ines-costa"),
    }
    assert all("old.csv" in r["reason"] for r in suppression.list())
    again = import_contacts(pid, people_file(tmp_path), "exclude", dry_run=False)
    assert again["suppressed"] == 0 and again["already_suppressed"] == 2
    assert len(suppression.list()) == 2


def test_do_not_contact_is_suppressed_in_signal_mode_and_never_stored_as_warm(tmp_path, pid):
    result = import_contacts(pid, people_file(tmp_path), "signal", dry_run=False)
    assert result["suppressed"] == 2 and "excluded" not in result
    assert count("contact_log") == 0
    with db.session() as conn:
        stored = db.query(conn, "SELECT first_name, relationship FROM known_contacts ORDER BY id")
    assert [(r["first_name"], r["relationship"]) for r in stored] == [
        ("Aria", "replied"), ("Omar", "contacted"), ("Hana", "customer"),
    ]  # fmt: skip


def test_signal_stores_known_contacts_and_flags_people_and_companies(ready_db, tmp_path):
    campaign = make_campaign()
    cid = campaign["id"]
    project_id = get_campaign(cid)["project_id"]
    company = add_company(cid)
    other_company = add_company(cid, "Contoso Freight", "contoso.example")
    quiet = add_company(cid, "Wingtip Air", "wingtip.example")
    aria = add_person(cid, company, "Aria", "Tan")
    omar = add_person(cid, other_company, "Omar", "Haddad", company="Contoso Freight")
    priya = add_person(cid, company, "Priya", "Nair")
    path = write_csv(
        tmp_path / "warm.csv",
        [
            row(first_name="aria", last_name="TAN", company="northwind logistics",
                last_contacted="2026-05-01", relationship="customer"),
            row(first_name="Someone", last_name="Else", company="Other Name",
                company_domain="contoso.example", email="someone@contoso.example"),
            row(first_name="Pat", last_name="Kim", company=COMPANY, linkedin_url="in/nope"),
        ],
    )  # fmt: skip
    result = import_contacts(project_id, path, "signal", dry_run=False)
    assert result["signal_new"] == 3 and result["signal_updated"] == 0
    assert count("contact_log") == 0 and count("known_contacts") == 3
    with db.session() as conn:
        stored = db.query(conn, "SELECT * FROM known_contacts ORDER BY id")
    assert stored[0]["key_name_company"] == "aria tan|northwind logistics"
    assert stored[0]["key_company"] == "northwind logistics"
    assert stored[0]["last_contacted"] == "2026-05-01"
    assert stored[0]["source"] == "import: warm.csv"

    flagged = contacts_import.known_for_people(project_id, [aria, omar, priya])
    assert flagged == {aria: {"relationship": "customer", "last_contacted": "2026-05-01"}}

    # companies: warm counts by name, and by domain when the file gave one
    listing = companies.list_companies(cid)["companies"]
    warm = {c["name"]: c.get("warm_contacts") for c in listing}
    assert warm == {COMPANY: 2, "Contoso Freight": 1, "Wingtip Air": None}
    assert quiet and "warm_contacts" not in next(c for c in listing if c["id"] == quiet)

    # the people list and the dashboard lists carry the flag
    listed = {p["id"]: p for p in people.list_people(cid)["people"]}
    assert listed[aria]["known_contact"]["relationship"] == "customer"
    assert "known_contact" not in listed[priya]
    leads = {p["id"]: p for p in browse.list_leads(cid)["leads"]}
    assert leads[aria]["known_contact"]["last_contacted"] == "2026-05-01"
    assert "known_contact" not in leads[omar]
    assert browse.get_person(aria)["person"]["known_contact"]["relationship"] == "customer"
    assert "known_contact" not in browse.get_person(priya)["person"]


def test_drafting_batch_carries_known_contact_only_for_imported_people(ready_db, tmp_path):
    campaign = make_campaign(channels=("linkedin",))
    cid = campaign["id"]
    project_id = get_campaign(cid)["project_id"]
    company = add_company(cid)
    aria = add_person(cid, company, "Aria", "Tan")
    priya = add_person(cid, company, "Priya", "Nair")
    path = write_csv(
        tmp_path / "known.csv",
        [row(first_name="Aria", last_name="Tan", company=COMPANY, relationship="replied",
             last_contacted="2026-04-02")],
    )  # fmt: skip

    def batch():
        reply = drafting.get_drafting_batch(cid, "linkedin_message")
        return {p["person_id"]: p for p in reply["people"]}

    assert "known_contact" not in batch()[aria]  # nothing imported yet
    import_contacts(project_id, path, "exclude", dry_run=False)
    assert "known_contact" not in batch()[aria]  # exclude only holds people back
    import_contacts(project_id, path, "signal", dry_run=False)
    got = batch()
    assert got[aria]["known_contact"] == {"relationship": "replied", "last_contacted": "2026-04-02"}
    assert "known_contact" not in got[priya]
    assert "known_contact" in (drafting.guide_text())


def test_known_contacts_are_per_project(ready_db, tmp_path, pid):
    campaign = make_campaign()
    cid = campaign["id"]
    company = add_company(cid)
    aria = add_person(cid, company, "Aria", "Tan")
    path = write_csv(
        tmp_path / "k.csv", [row(first_name="Aria", last_name="Tan", company=COMPANY)]
    )
    import_contacts(pid, path, "signal", dry_run=False)  # a different project
    assert contacts_import.known_for_people(get_campaign(cid)["project_id"], [aria]) == {}
    assert contacts_import.known_for_people(None, [aria]) == {}
    assert contacts_import.known_for_people(pid, []) == {}


def test_reimporting_the_same_file_creates_no_duplicates(tmp_path, pid):
    path = people_file(tmp_path)
    first = import_contacts(pid, path, "both", dry_run=False)
    after_first = table_counts()
    second = import_contacts(pid, path, "both", dry_run=False)
    assert table_counts() == after_first
    assert first["excluded"] == 5 and second["excluded"] == 0 and second["already_logged"] == 5
    assert first["signal_new"] == 3 and second["signal_new"] == 0 and second["signal_updated"] == 0
    assert second["suppressed"] == 0
    # a dry run of the same file says the same
    third = import_contacts(pid, path, "both")
    assert third["would_exclude"] == 0 and third["would_signal"] == 0
    assert table_counts() == after_first


def test_a_newer_last_contacted_on_reimport_holds_the_person_back(tmp_path, pid):
    aria = dict(first_name="Aria", last_name="Tan", email="aria@northwind.example", company=COMPANY)
    older = write_csv(tmp_path / "v1.csv", [row(**aria, last_contacted=days_ago(400).isoformat())])
    newer = write_csv(tmp_path / "v2.csv", [row(**aria, last_contacted=days_ago(3).isoformat())])
    person = {k: aria[k] for k in ("first_name", "last_name", "email", "company")}
    assert import_contacts(pid, older, "exclude", dry_run=False)["excluded"] == 1
    assert projects.recently_contacted(pid, person) is None  # 400 days is outside the cooldown
    second = import_contacts(pid, newer, "exclude", dry_run=False)
    assert second["excluded"] == 1 and second["already_logged"] == 0
    hit = projects.recently_contacted(pid, person)
    assert hit is not None and hit["channel"] == "imported" and hit["at"].startswith(
        days_ago(3).isoformat()
    )
    assert count("contact_log") == 1  # moved forward, not repeated
    # the identical newer file again, and the older one again, change nothing
    again = import_contacts(pid, newer, "exclude", dry_run=False)
    assert again["excluded"] == 0 and again["already_logged"] == 1
    back = import_contacts(pid, older, "exclude", dry_run=False)
    assert back["excluded"] == 0 and back["already_logged"] == 1
    assert projects.recently_contacted(pid, person)["at"] == hit["at"]
    assert count("contact_log") == 1


def test_an_undated_contact_is_not_refreshed_by_a_later_import(pid):
    person = {"first_name": "Omar", "last_name": "Reyes", "company": COMPANY}
    with db.session() as conn:
        projects.log_contacts_in(conn, pid, [person], "imported")
        conn.execute("UPDATE contact_log SET at = '2026-01-01T00:00:00Z'")
        result = projects.log_contacts_in(conn, pid, [person], "imported")
        assert result == {"logged": 0, "skipped": 1}
        assert conn.execute("SELECT at FROM contact_log").fetchone()[0] == "2026-01-01T00:00:00Z"


def test_the_duplicate_check_uses_a_key_index_for_import_rows(pid):
    keys = {
        "key_linkedin": "in/aria-tan",
        "key_email": "aria@northwind.example",
        "key_name_company": "aria tan|northwind logistics",
    }
    cases = [
        (keys, "idx_contact_log_linkedin"),
        ({**keys, "key_linkedin": ""}, "idx_contact_log_email"),
        ({**keys, "key_linkedin": "", "key_email": ""}, "idx_contact_log_name"),
    ]
    with db.session() as conn:
        for entry, index in cases:
            sql, first = projects._same_log_sql(entry)
            params = (pid, first, None, "imported", None, entry["key_linkedin"],
                      entry["key_email"], entry["key_name_company"])  # fmt: skip
            plan = " ".join(r[3] for r in conn.execute("EXPLAIN QUERY PLAN " + sql, params))
            assert index in plan and "idx_contact_log_person" not in plan


def test_a_large_import_logs_every_row_once(tmp_path, pid):
    rows = [
        row(first_name=f"Person{i}", last_name="Tan", email=f"p{i}@northwind.example",
            company=COMPANY, last_contacted=days_ago(5).isoformat())
        for i in range(600)
    ]  # fmt: skip
    path = write_csv(tmp_path / "big.csv", rows)
    assert import_contacts(pid, path, "exclude", dry_run=False)["excluded"] == 600
    assert import_contacts(pid, path, "exclude", dry_run=False)["already_logged"] == 600
    assert count("contact_log") == 600


def test_a_changed_file_updates_known_contacts_instead_of_adding(tmp_path, pid):
    path = write_csv(
        tmp_path / "v1.csv",
        [row(first_name="Aria", last_name="Tan", company=COMPANY, last_contacted="2026-03-01")],
    )
    import_contacts(pid, path, "signal", dry_run=False)
    path2 = write_csv(
        tmp_path / "v2.csv",
        [
            row(first_name="Aria", last_name="Tan", company=COMPANY, email="aria@northwind.example",
                title="Head of Ops", relationship="customer", last_contacted="2026-09-01"),
        ],
    )  # fmt: skip
    result = import_contacts(pid, path2, "signal", dry_run=False)
    assert result["signal_new"] == 0 and result["signal_updated"] == 1
    with db.session() as conn:
        stored = db.query(conn, "SELECT * FROM known_contacts")
    assert len(stored) == 1
    only = stored[0]
    assert only["key_email"] == "aria@northwind.example" and only["title"] == "Head of Ops"
    assert only["relationship"] == "customer" and only["last_contacted"] == "2026-09-01"
    # an older file never takes anything back
    import_contacts(pid, path, "signal", dry_run=False)
    with db.session() as conn:
        again = db.query(conn, "SELECT * FROM known_contacts")[0]
    assert again["relationship"] == "customer" and again["last_contacted"] == "2026-09-01"


def test_errors_for_unusable_requests(tmp_path, pid):
    path = people_file(tmp_path)
    with pytest.raises(ContactsImportError) as bad_mode:
        import_contacts(pid, path, "everything")
    assert bad_mode.value.code == "invalid_mode"
    with pytest.raises(ContactsImportError) as missing:
        import_contacts(pid, tmp_path / "nope.csv", "exclude")
    assert missing.value.code == "unreadable_file"
    assert missing.value.as_response()["error"]["hint"]
    with pytest.raises(projects.ProjectError) as unknown:
        import_contacts(4242, path, "exclude")
    assert unknown.value.code == "not_found"
    inside = tmp_path / "inside"
    inside.mkdir()
    with pytest.raises(ContactsImportError) as outside:
        import_contacts(pid, path, "exclude", file_root=inside)
    assert outside.value.code == "file_outside_workspace"
    copy = shutil.copy(path, inside / "old.csv")
    assert import_contacts(pid, copy, "exclude", file_root=inside)["ok"]


def test_log_contacts_accepts_a_date_per_row(pid):
    result = projects.log_contacts(
        pid,
        [
            {"email": "a@x.example", "at": "2026-01-02"},
            {"email": "b@x.example"},
        ],
        "imported",
        at="2026-02-03",
    )
    assert result == {"logged": 2, "skipped": 0}
    with db.session() as conn:
        days = {r["key_email"]: r["at"][:10] for r in db.query(conn, "SELECT * FROM contact_log")}
    assert days == {"a@x.example": "2026-01-02", "b@x.example": "2026-02-03"}
    with pytest.raises(projects.ProjectError):
        projects.log_contacts(pid, [{"email": "c@x.example", "at": "soon"}], "imported")
    assert count("contact_log") == 2


# --- schema v3 --------------------------------------------------------------------------------

V2_TABLES = (
    "offers", "projects", "campaigns", "companies", "people", "contact_log", "suppression",
    "settings",
)  # fmt: skip


def build_v2_database(path, monkeypatch):
    with monkeypatch.context() as patch:
        patch.setattr(db, "MIGRATIONS", db.MIGRATIONS[:1])
        assert db.init_db(path) == 2
    with db.session(path) as conn:
        offer_id = db.insert(conn, "offers", name="Northwind Training", source_type="text")
        project_id = db.insert(conn, "projects", name="Northwind Training", offer_id=offer_id)
        campaign_id = db.insert(conn, "campaigns", name="SG", offer_id=offer_id,
                                project_id=project_id)  # fmt: skip
        company = db.insert(
            conn, "companies", campaign_id=campaign_id, name=COMPANY, status="approved"
        )
        db.insert(conn, "people", campaign_id=campaign_id, company_id=company, first_name="Aria",
                  last_name="Tan", raw_name="Aria Tan", status="confirmed")  # fmt: skip
        db.insert(conn, "contact_log", project_id=project_id, key_email="aria@northwind.example",
                  channel="imported", at="2026-09-01T00:00:00.000Z")  # fmt: skip
        db.insert(conn, "suppression", value="nobody@example.com", kind="email")
        db.insert(conn, "settings", key="k", value="v")
    return project_id


def test_migration_v3_on_a_copy_of_a_v2_database(tmp_path, monkeypatch):
    original = tmp_path / "v2.db"
    project_id = build_v2_database(original, monkeypatch)
    copy = tmp_path / "copy" / "leadgen.db"
    copy.parent.mkdir()
    shutil.copy2(original, copy)

    def counts(path):
        with db.session(path) as conn:
            return {t: conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in V2_TABLES}

    before = counts(copy)
    with monkeypatch.context() as patch:  # stop at v3 (v4 has its own test, T25)
        patch.setattr(db, "MIGRATIONS", db.MIGRATIONS[:2])
        assert db.init_db(copy) == 3
    assert counts(copy) == before  # nothing in the v2 tables changed
    with db.session(original) as conn:
        assert db.current_version(conn) == 2  # the original file was not touched
        assert "known_contacts" not in table_names(conn)
    with db.session(copy) as conn:
        assert db.current_version(conn) == 3
        assert conn.execute("SELECT COUNT(*) FROM known_contacts").fetchone()[0] == 0
        indexes = {r[1] for r in conn.execute("PRAGMA index_list(known_contacts)")}
        assert {"idx_known_linkedin", "idx_known_email", "idx_known_name",
                "idx_known_company"} <= indexes  # fmt: skip
        assert conn.execute("PRAGMA foreign_key_check").fetchall() == []
        # a known contact can be stored, and goes with its project
        db.insert(conn, "known_contacts", project_id=project_id, key_email="a@x.example",
                  imported_at="2026-09-01", updated_at="2026-09-01")  # fmt: skip
        conn.execute("DELETE FROM contact_log")
        conn.execute("DELETE FROM campaigns")
        conn.execute("DELETE FROM projects")
        assert conn.execute("SELECT COUNT(*) FROM known_contacts").fetchone()[0] == 0
    # running it again changes nothing
    with monkeypatch.context() as patch:
        patch.setattr(db, "MIGRATIONS", db.MIGRATIONS[:2])
        assert db.init_db(copy) == 3


def test_failed_v3_migration_rolls_back(tmp_path, monkeypatch):
    path = tmp_path / "v2.db"
    build_v2_database(path, monkeypatch)
    monkeypatch.setattr(
        db, "MIGRATIONS",
        [db.MIGRATIONS[0], (3, db._V3_SQL + "\nINSERT INTO no_such_table VALUES (1);")],
    )  # fmt: skip
    with pytest.raises(sqlite3.Error):
        db.init_db(path)
    with db.session(path) as conn:
        assert db.current_version(conn) == 2
        assert "known_contacts" not in table_names(conn)


def test_the_current_schema_has_known_contacts():
    assert db.SCHEMA_VERSION >= 3 and "known_contacts" in db.TABLES


# --- command line -----------------------------------------------------------------------------


def test_cli_template_and_import(tmp_path, capsys, workspace, pid):
    from leadgen.cli import main

    out = tmp_path / "mine.csv"
    assert main(["contacts", "template", "--out", str(out)]) == 0
    assert out.exists() and str(out) in capsys.readouterr().out
    xlsx = tmp_path / "mine.xlsx"
    assert main(["contacts", "template", "--xlsx", "--out", str(xlsx)]) == 0
    assert openpyxl.load_workbook(xlsx).sheetnames[0] == "Contacts"
    assert main(["contacts", "template"]) == 0  # default place: the workspace templates folder
    assert (workspace / "templates" / "contacts_template.csv").exists()
    capsys.readouterr()

    path = people_file(tmp_path)
    before = table_counts()
    assert main(["contacts", "import", str(path), "--project", str(pid), "--mode", "both"]) == 0
    shown = capsys.readouterr().out
    assert "Preview only" in shown and "--apply" in shown
    assert table_counts() == before
    assert main(["contacts", "import", str(path), "--project", str(pid), "--apply"]) == 0
    assert "Preview only" not in capsys.readouterr().out
    assert count("contact_log") == 5 and count("known_contacts") == 0  # default mode is exclude
    assert main(["contacts", "import", str(tmp_path / "no.csv"), "--project", str(pid)]) == 1
    assert "no file" in capsys.readouterr().out.lower()
    assert main(["contacts", "import", str(path), "--project", "999"]) == 1
    capsys.readouterr()
    messy = write_csv(tmp_path / "messy.csv", [row(first_name="A", last_name="B", email="bad",
                                                   linkedin_url="http://x.example")])  # fmt: skip
    assert main(["contacts", "import", str(messy), "--project", str(pid)]) == 0
    assert "Problems" in capsys.readouterr().out
