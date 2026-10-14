"""Offer service: ingest results, "to confirm" detection, chunked reading and the profile."""

from __future__ import annotations

import copy

import httpx
import pytest
import respx
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from leadgen import db
from leadgen.services import offer, offer_web
from leadgen.services.offer import OfferError, OfferProfile

pytestmark = pytest.mark.usefixtures("ready_db")

BASE = "https://www.northwind.example.com"

PROFILE = {
    "name": "Northwind Safety Training",
    "one_liner": "Practical safety training for warehouse teams.",
    "description": "Northwind runs on-site safety courses. Courses are led by certified trainers.",
    "services": [
        {"name": "Forklift course", "what": "One-day certification", "best_for": ["Warehouses"]}
    ],
    "audiences": ["Heads of operations", "L&D managers"],
    "confirmed_claims": ["Courses are led by certified trainers"],
    "forbidden_claims": ["98% pass rate"],
    "proof_points": [{"claim": "Trained 40 teams in 2025", "source": "example.com/about"}],
    "ctas": ["15-minute call"],
    "sender": {"name": "Aria Tan", "company": "Northwind Logistics", "role": "Account Director"},
    "tone": "professional, plain, no hype",
    "language": "en-GB",
}


def profile(**changes):
    data = copy.deepcopy(PROFILE)
    data.update(changes)
    return data


def raw_text(offer_id: int) -> str:
    with db.session() as conn:
        return db.get(conn, "offers", offer_id)["raw_text"]


@pytest.fixture
def offer_id():
    text = "Northwind sells safety training.\nCourses are led by certified trainers."
    return offer.ingest_text(text)["offer_id"]


# --- ingest_text / ingest_file ------------------------------------------------------------------


def test_ingest_text():
    result = offer.ingest_text("  Northwind sells safety training.\r\n\r\n\r\n\r\nCall us.  ")
    assert result["ok"] and result["source_type"] == "text"
    assert result["name"] == "Northwind sells safety training."
    assert raw_text(result["offer_id"]) == "Northwind sells safety training.\n\nCall us."
    assert result["total_chars"] == len(raw_text(result["offer_id"]))
    assert result["unconfirmed_snippets"] == [] and result["unconfirmed_count"] == 0
    assert offer.ingest_text("x", name="  My   offer ")["name"] == "My offer"


def test_ingest_text_empty():
    result = offer.ingest_text("  \n ")
    assert result["ok"] is False and result["error"]["code"] == "empty_text"
    assert result["error"]["hint"]


def test_ingest_text_is_capped(monkeypatch):
    monkeypatch.setattr(offer, "MAX_RAW_CHARS", 20)
    result = offer.ingest_text("word " * 20)
    assert result["total_chars"] == 20
    assert any("cut" in w for w in result["warnings"])


def test_ingest_file_txt_and_errors(tmp_path):
    path = tmp_path / "Northwind Offer.md"
    path.write_text("# Offer\nWe train teams.\nPrice: TBC", encoding="utf-8")
    result = offer.ingest_file(path)
    assert result["ok"] and result["source_type"] == "file" and result["name"] == "Northwind Offer"
    assert result["source_ref"] == str(path.resolve())
    assert result["unconfirmed_snippets"] == ["Price: TBC"]

    bad = offer.ingest_file(tmp_path / "data.xlsx")
    assert bad["ok"] is False and bad["error"]["code"] in ("file_not_found", "unsupported_type")
    (tmp_path / "data.xlsx").write_bytes(b"x")
    bad = offer.ingest_file(tmp_path / "data.xlsx")
    assert bad["error"]["code"] == "unsupported_type" and ".pdf" in bad["error"]["message"]


def test_ingest_docx_reports_highlighted_placeholders(tmp_path):
    doc = Document()
    doc.add_paragraph("We run courses.")
    para = doc.add_paragraph("Pass rate: ")
    para.add_run("95%").font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.save(str(tmp_path / "offer.docx"))
    result = offer.ingest_file(tmp_path / "offer.docx")
    assert result["unconfirmed_snippets"] == ["Pass rate: [HIGHLIGHT: 95%]"]


# --- unconfirmed detection ----------------------------------------------------------------------


def test_find_unconfirmed_markers():
    text = "\n".join(
        [
            "# https://example.com/about",
            "Clients include [CONFIRM: Acme] and others",
            "Case study [PROVIDE link]",
            "Price tbc next week",
            "TODO add testimonials",
            "Headcount to be confirmed",
            "Savings of XX% per year",
            "Fee: GBP XXk",
            "TBD",
            "Hello {{first_name}}",
            "<insert logo>",
            "Normal sentence about outcomes.",
            "Size XXL shirts and ABCTBCD and todos",
            "Clients include [CONFIRM: Acme] and others",
        ]
    )
    snippets, total = offer.find_unconfirmed(text)
    assert snippets == [
        "Clients include [CONFIRM: Acme] and others",
        "Case study [PROVIDE link]",
        "Price tbc next week",
        "TODO add testimonials",
        "Headcount to be confirmed",
        "Savings of XX% per year",
        "Fee: GBP XXk",
        "TBD",
        "Hello {{first_name}}",
        "<insert logo>",
    ]
    assert total == 11  # the duplicate line counts but is listed once


def test_find_unconfirmed_long_line_window_and_limit():
    line = "a" * 400 + " [CONFIRM: figure] " + "b" * 400
    (snippet,), total = offer.find_unconfirmed(line)
    assert "[CONFIRM: figure]" in snippet and len(snippet) <= offer.SNIPPET_CHARS and total == 1
    many = "\n".join(f"item {i} TBC" for i in range(80))
    snippets, total = offer.find_unconfirmed(many, limit=10)
    assert len(snippets) == 10 and total == 80


def test_ingest_text_returns_unconfirmed_snippets():
    result = offer.ingest_text("We help teams.\nResults: [CONFIRM: 30% faster]")
    assert result["unconfirmed_snippets"] == ["Results: [CONFIRM: 30% faster]"]
    assert result["unconfirmed_count"] == 1


# --- ingest_url ---------------------------------------------------------------------------------


@pytest.fixture
def fast_crawl(monkeypatch):
    monkeypatch.setattr(offer_web, "REQUESTS_PER_SECOND", 10_000.0)


def html(body: str) -> str:
    return f"<html><head><title>Northwind</title></head><body>{body}</body></html>"


@respx.mock
def test_ingest_url_builds_page_blocks(fast_crawl):
    respx.get(f"{BASE}/robots.txt").mock(return_value=httpx.Response(404))
    respx.get(f"{BASE}/").mock(
        return_value=httpx.Response(
            200, html=html('<h1>Safety</h1><p>Training. Rate: TBC</p><a href="/about">About</a>')
        )
    )
    respx.get(f"{BASE}/about").mock(
        return_value=httpx.Response(200, html=html("<p>About Northwind</p>"))
    )
    respx.get(url__regex=rf"^{BASE}/.+").mock(return_value=httpx.Response(404))
    result = offer.ingest_url("www.northwind.example.com")
    assert result["ok"] and result["name"] == "northwind.example.com"
    assert [p["url"] for p in result["pages"]] == [f"{BASE}/", f"{BASE}/about"]
    text = raw_text(result["offer_id"])
    assert text.startswith(f"# {BASE}/\n## Northwind\n## Safety\nTraining. Rate: TBC")
    assert f"\n\n# {BASE}/about\n" in text
    assert result["unconfirmed_snippets"] == ["Training. Rate: TBC"]
    assert result["source_type"] == "url" and result["source_ref"] == "https://www.northwind.example.com"


@respx.mock
def test_ingest_url_404_returns_structured_error(fast_crawl):
    respx.get(f"{BASE}/robots.txt").mock(return_value=httpx.Response(404))
    respx.get(f"{BASE}/").mock(return_value=httpx.Response(404, text="gone"))
    result = offer.ingest_url(BASE + "/")
    assert result["ok"] is False
    error = result["error"]
    assert error["code"] == "http_error" and "404" in error["message"]
    assert "paste" in error["hint"] and "file" in error["hint"]
    assert offer.list_offers() == []  # nothing stored


@respx.mock
def test_ingest_url_no_text_and_bad_url(fast_crawl):
    respx.get(f"{BASE}/robots.txt").mock(return_value=httpx.Response(404))
    respx.get(f"{BASE}/").mock(return_value=httpx.Response(200, html="<html><body></body></html>"))
    assert offer.ingest_url(BASE)["error"]["code"] == "no_text"
    result = offer.ingest_url("ftp://example.com")
    assert result["error"]["code"] == "invalid_url" and "paste" in result["error"]["hint"]


@respx.mock
def test_ingest_url_google_doc():
    doc_url = "https://docs.google.com/document/d/abc123/edit?usp=sharing"
    export = "https://docs.google.com/document/d/abc123/export?format=txt"
    route = respx.get(export).mock(
        return_value=httpx.Response(
            200, text="﻿Northwind Offer\nPrice TBC", headers={"content-type": "text/plain"}
        )
    )
    result = offer.ingest_url(doc_url)
    assert route.called and result["ok"] and result["name"] == "Northwind Offer"
    assert raw_text(result["offer_id"]) == f"# {doc_url}\nNorthwind Offer\nPrice TBC"
    assert result["unconfirmed_snippets"] == ["Price TBC"]
    assert result["pages"] == [{"url": doc_url, "chars": len("Northwind Offer\nPrice TBC")}]


@respx.mock
def test_ingest_url_private_google_doc():
    export = "https://docs.google.com/document/d/abc123/export?format=txt"
    respx.get(export).mock(return_value=httpx.Response(403, html="<html>no</html>"))
    result = offer.ingest_url("https://docs.google.com/document/d/abc123/edit")
    assert result["ok"] is False and result["error"]["code"] == "private_document"
    assert ".docx" in result["error"]["hint"] and "paste" in result["error"]["hint"]
    sheet = offer.ingest_url("https://docs.google.com/spreadsheets/d/abc/edit")
    assert sheet["error"]["code"] == "unsupported_google_link"


# --- reading in chunks --------------------------------------------------------------------------


def walk(offer_id: int, max_chars: int) -> list[str]:
    chunks, cursor = [], 0
    while cursor is not None:
        part = offer.read_offer_text(offer_id, cursor, max_chars)
        assert 0 < len(part["text"]) <= max_chars
        chunks.append(part["text"])
        cursor = part["next_cursor"]
    return chunks


def test_cursor_walks_the_full_text_without_loss():
    text = "\n".join(f"Line {i} " + "word " * (i % 17) for i in range(12))
    offer_id = offer.ingest_text(text)["offer_id"]
    stored = raw_text(offer_id)
    for size in (3, 50, 40_000):
        assert "".join(walk(offer_id, size)) == stored


def test_chunks_prefer_line_breaks_and_text_without_newlines():
    offer_id = offer.ingest_text("alpha beta\ngamma delta\nepsilon zeta")["offer_id"]
    first = offer.read_offer_text(offer_id, 0, 15)
    assert first["text"] == "alpha beta\n" and first["next_cursor"] == 11
    assert first["total_chars"] == len(raw_text(offer_id))
    solid = offer.ingest_text("x" * 250)["offer_id"]
    assert "".join(walk(solid, 100)) == "x" * 250


def test_read_offer_text_limits_and_errors(offer_id):
    whole = offer.read_offer_text(offer_id)
    assert whole["next_cursor"] is None and whole["text"] == raw_text(offer_id)
    past = offer.read_offer_text(offer_id, cursor=10_000)
    assert past == {"text": "", "next_cursor": None, "total_chars": whole["total_chars"]}
    big = offer.ingest_text("y" * 100_000)["offer_id"]
    assert len(offer.read_offer_text(big, 0, 10**9)["text"]) == offer.MAX_CHUNK_CHARS
    with pytest.raises(OfferError) as info:
        offer.read_offer_text(offer_id, cursor=-1)
    assert info.value.code == "invalid_cursor"
    with pytest.raises(OfferError) as info:
        offer.read_offer_text(9999)
    assert info.value.code == "not_found"
    assert info.value.as_response()["error"]["code"] == "not_found"


# --- get / list ---------------------------------------------------------------------------------


def test_get_and_list_offers(offer_id):
    second = offer.ingest_text("Second offer", name="Second")["offer_id"]
    assert [o["id"] for o in offer.list_offers()] == [second, offer_id]
    details = offer.get_offer(offer_id)
    assert details["profile"] is None and details["has_profile"] is False
    assert "raw_text" not in details and details["total_chars"] == len(raw_text(offer_id))
    with pytest.raises(OfferError):
        offer.get_offer(12345)


# --- profile ------------------------------------------------------------------------------------


def test_save_profile_round_trip(offer_id):
    saved = offer.save_profile(offer_id, PROFILE)
    assert isinstance(saved, OfferProfile)
    stored = offer.get_offer(offer_id)
    assert stored["has_profile"] and stored["name"] == "Northwind Safety Training"
    assert stored["profile"] == saved.model_dump(mode="json")
    assert OfferProfile.model_validate(stored["profile"]) == saved
    assert saved.sender.email is None
    assert offer.list_offers()[0]["has_profile"] is True


def test_profile_is_normalised_and_deduped(offer_id):
    messy = profile(
        name="  Northwind   Safety\nTraining ",
        audiences=["Heads of ops", "heads  of OPS", "", "  L&D "],
        confirmed_claims=["Trained  by experts", "trained by experts"],
        forbidden_claims=[],
        services=[
            {"name": "Forklift", "what": " One   day ", "best_for": ["A", "a", " "]},
            {"name": "forklift", "what": "Duplicate", "best_for": []},
        ],
        proof_points=[
            {"claim": "40 teams", "source": "about page"},
            {"claim": "40  TEAMS", "source": "other"},
        ],
        sender={"name": "Aria Tan", "company": "Northwind", "role": "Director", "email": "  ",
                "website": " https://northwind.example.com "},
        language="en_gb",
    )
    saved = offer.save_profile(offer_id, messy)
    assert saved.name == "Northwind Safety Training"
    assert saved.audiences == ["Heads of ops", "L&D"]
    assert saved.confirmed_claims == ["Trained by experts"]
    assert [s.name for s in saved.services] == ["Forklift"]
    assert saved.services[0].what == "One day" and saved.services[0].best_for == ["A"]
    assert len(saved.proof_points) == 1 and saved.proof_points[0].source == "about page"
    assert saved.sender.email is None and saved.sender.website == "https://northwind.example.com"
    assert saved.language == "en-GB"


def test_missing_required_fields(offer_id):
    data = profile()
    del data["name"], data["language"]
    del data["sender"]["company"]
    with pytest.raises(OfferError) as info:
        offer.save_profile(offer_id, data)
    error = info.value
    assert error.code == "invalid_profile"
    for field in ("name", "language", "sender.company"):
        assert field in error.message
    assert error.hint
    assert offer.get_offer(offer_id)["profile"] is None  # nothing saved


def test_invalid_values_do_not_echo_input(offer_id):
    with pytest.raises(OfferError) as info:
        offer.save_profile(offer_id, profile(one_liner="   ", tone=42, language="not a locale"))
    message = info.value.message
    assert "one_liner" in message and "tone" in message and "language" in message
    assert "not a locale" not in message


def test_unknown_and_incomplete_nested_fields(offer_id):
    for bad in (
        profile(surprise="x"),
        profile(proof_points=[{"claim": "40 teams", "source": " "}]),
        profile(services=[{"name": "A"}]),
        profile(sender="Aria Tan"),
    ):
        with pytest.raises(OfferError) as info:
            offer.save_profile(offer_id, bad)
        assert info.value.code == "invalid_profile"


@pytest.mark.parametrize(
    "forbidden",
    [["courses are led by certified trainers"], ["  Courses are   led by certified trainers. "]],
)
def test_overlapping_claims_are_rejected(offer_id, forbidden):
    with pytest.raises(OfferError) as info:
        offer.save_profile(offer_id, profile(forbidden_claims=forbidden))
    assert info.value.code == "invalid_profile"
    assert "overlap" in info.value.message and "certified trainers" in info.value.message
    assert offer.get_offer(offer_id)["profile"] is None


def test_save_profile_unknown_offer():
    with pytest.raises(OfferError) as info:
        offer.save_profile(777, PROFILE)
    assert info.value.code == "not_found"


def test_resaving_replaces_profile(offer_id):
    offer.save_profile(offer_id, PROFILE)
    offer.save_profile(offer_id, profile(tone="warm, direct"))
    assert offer.get_offer(offer_id)["profile"]["tone"] == "warm, direct"


# --- guide --------------------------------------------------------------------------------------


def test_guide_text_covers_the_rules():
    guide = offer.guide_text()
    needles = ("forbidden_claims", "to confirm", "proof_points", "[HIGHLIGHT", "confirm with")
    for needle in needles:
        assert needle in guide
