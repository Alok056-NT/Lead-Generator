"""Import contacts the owner already has (CSV, Excel, LinkedIn's Connections.csv) (T24).

Each file is read into an `ImportPreview` first (`read_contacts`: it never raises on bad rows, it
lists them), then `import_contacts` uses it in one of three modes:

- `exclude`: every contact is written to the project's contact log (channel `imported`, at the
  `last_contacted` date), so L24 holds them back during the project's cooldown;
- `signal`: contacts are stored in `known_contacts`, so people are flagged `known_contact` and
  companies count as warm (`warm_contacts`);
- `both`.

Contacts whose relationship is `do_not_contact` go on the global do-not-contact list in every
mode (an address or LinkedIn link; a name-only row cannot be listed, a warning says so) and are
never stored as warm contacts. A dry run does all the work inside a transaction that is rolled
back, so its counts are exact and nothing is written.

Template and import columns: `COLUMNS`. Callers must run `db.init_db()` first.
"""

from __future__ import annotations

import csv
import io
import re
import sqlite3
from collections.abc import Callable, Iterable, Iterator, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, timedelta
from importlib import resources
from pathlib import Path
from typing import Any

from leadgen import db
from leadgen.core import paths
from leadgen.core.text import normalize_whitespace
from leadgen.services import projects, suppression
from leadgen.services.confirm import canonical_url, name_key

COLUMNS = (
    "first_name", "last_name", "full_name", "email", "linkedin_url", "company",
    "company_domain", "title", "country", "last_contacted", "channel", "relationship", "notes",
)  # fmt: skip
MODES = ("exclude", "signal", "both")
RELATIONSHIPS = ("contacted", "replied", "customer", "partner", "do_not_contact")
CHANNELS = ("email", "linkedin", "both")
MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 50_000
MAX_STORED_MESSAGES = 1000  # errors and warnings kept in the preview (all are counted)
SHOWN = 20  # errors shown in a reply
HEADER_SCAN_ROWS = 60
MAX_BLANK_RUN = 2000  # an Excel sheet can report a million formatted empty rows
MAX_SCANNED_ROWS = MAX_ROWS * 4
SHEET_NAME = "Contacts"
GUIDE_SHEET_NAME = "How to fill this in"
TEMPLATE_CSV = "contacts_template.csv"
TEMPLATE_XLSX = "contacts_template.xlsx"
_LIMITS = {
    "first_name": 120, "last_name": 120, "full_name": 240, "company": 200, "title": 200,
    "country": 80, "notes": 500,
}  # fmt: skip

# What each column means: (required text, explanation, example). One line each in the Excel guide.
COLUMN_GUIDE: dict[str, tuple[str, str, str]] = {
    "first_name": ("one of these", "First name. Leave empty if you fill full_name.", "Aria"),
    "last_name": ("one of these", "Last name. Leave empty if you fill full_name.", "Tan"),
    "full_name": (
        "one of these", "Whole name; used only when first_name and last_name are empty.",
        "Aria Tan",
    ),
    "email": ("recommended", "Work or personal email address.", "aria@northwind.example"),
    "linkedin_url": (
        "recommended", "Link to the person's LinkedIn profile (linkedin.com/in/...).",
        "https://www.linkedin.com/in/aria-tan",
    ),
    "company": (
        "recommended", "Company name. With a name it lets us match people who have no email.",
        "Northwind Logistics",
    ),
    "company_domain": ("optional", "Company website address.", "northwind.example"),
    "title": ("optional", "Job title.", "Head of Operations"),
    "country": ("optional", "Country.", "Singapore"),
    "last_contacted": (
        "optional",
        "When you last contacted them: 2026-09-30, 30/09/2026 or 30-09-2026 (or an Excel date). "
        "Empty means today.",
        "2026-09-30",
    ),
    "channel": ("optional", "How you contacted them: email, linkedin or both.", "email"),
    "relationship": (
        "optional",
        "contacted (default), replied, customer, partner or do_not_contact. "
        "do_not_contact is never contacted.",
        "customer",
    ),
    "notes": ("optional", "Anything you want to remember.", "Met at the 2026 expo"),
}
_NEEDS = (
    "Each row needs an email, or a LinkedIn link, or a name together with a company. "
    "Column order does not matter and extra columns are ignored."
)

# header spelling (lower case, letters and digits only joined by `_`) -> column
_ALIASES: dict[str, str] = {}
for _column, _names in {
    "first_name": ("first_name", "firstname", "first", "given_name"),
    "last_name": ("last_name", "lastname", "last", "surname", "family_name"),
    "full_name": ("full_name", "fullname", "name", "contact_name"),
    "email": ("email", "email_address", "e_mail", "mail"),
    "linkedin_url": (
        "linkedin_url", "linkedin", "linkedin_profile", "linkedin_profile_url", "profile_url",
        "url",
    ),
    "company": ("company", "company_name", "organization", "organisation", "employer"),
    "company_domain": ("company_domain", "domain", "website", "company_website"),
    "title": ("title", "job_title", "position", "role"),
    "country": ("country",),
    "last_contacted": ("last_contacted", "last_contact", "last_contacted_on", "date_contacted"),
    "channel": ("channel",),
    "relationship": ("relationship", "status"),
    "notes": ("notes", "note", "comments"),
    "connected_on": ("connected_on",),  # LinkedIn Connections.csv
}.items():
    for _name in _names:
        _ALIASES[_name] = _column
_IDENTITY = ("first_name", "last_name", "full_name", "email", "linkedin_url", "company")
_RELATIONSHIP_ALIASES = {
    "dnc": "do_not_contact", "do_not_email": "do_not_contact", "unsubscribed": "do_not_contact",
    "opted_out": "do_not_contact", "opt_out": "do_not_contact", "optout": "do_not_contact",
    "do_not_contact": "do_not_contact", "never": "do_not_contact",
    "client": "customer", "contact": "contacted", "reply": "replied", "responded": "replied",
}  # fmt: skip
_CHANNEL_ALIASES = {"e_mail": "email", "mail": "email", "linked_in": "linkedin", "li": "linkedin"}
_RANK = {name: i for i, name in enumerate(RELATIONSHIPS)}
_STRONGEST = ("contacted", "replied", "partner", "customer", "do_not_contact")  # weak to strong
_DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d %b %Y", "%d %B %Y")
_BAD_CELL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_MIN_YEAR = 1990


class ContactsImportError(Exception):
    """An import request failed. `code` is stable text; `hint` tells the AI what to do."""

    def __init__(self, code: str, message: str, hint: str = "") -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.hint = hint

    def as_response(self) -> dict[str, Any]:
        """The `{"ok": False, "error": {...}}` shape used by every tool."""
        return {
            "ok": False,
            "error": {"code": self.code, "message": self.message, "hint": self.hint},
        }


@dataclass
class Contact:
    """One person after cleaning (rows of the same person merged)."""

    first_name: str = ""
    last_name: str = ""
    email: str = ""
    linkedin_url: str = ""
    company: str = ""
    company_domain: str = ""
    title: str = ""
    country: str = ""
    last_contacted: str = ""  # YYYY-MM-DD or ""
    channel: str = ""
    relationship: str = "contacted"
    notes: str = ""
    rows: list[int] = field(default_factory=list)

    def keys(self) -> dict[str, str]:
        """The three contact-memory keys (see `projects.make_keys`)."""
        return projects.make_keys(self.as_person())

    def as_person(self) -> dict[str, str]:
        """The fields `projects.make_keys` reads."""
        return {
            "first_name": self.first_name, "last_name": self.last_name, "email": self.email,
            "linkedin_url": self.linkedin_url, "company": self.company,
        }  # fmt: skip

    @property
    def name(self) -> str:
        return normalize_whitespace(f"{self.first_name} {self.last_name}")


@dataclass
class ImportPreview:
    """What a file holds, before anything is written. `error` is set when it cannot be read."""

    path: str = ""
    format: str = ""  # csv / xlsx
    source: str = "template"  # template / linkedin_connections
    error: str = ""
    rows_read: int = 0
    valid_rows: int = 0
    rows_skipped: int = 0
    duplicates: int = 0
    error_count: int = 0
    warning_count: int = 0
    errors: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[dict[str, Any]] = field(default_factory=list)
    unknown_columns: list[str] = field(default_factory=list)
    contacts: list[Contact] = field(default_factory=list)

    def _add(self, store: list[dict[str, Any]], row: int, column: str, problem: str) -> None:
        if len(store) < MAX_STORED_MESSAGES:
            store.append({"row": row, "column": column, "problem": problem})

    def add_error(self, row: int, column: str, problem: str) -> None:
        self.error_count += 1
        self._add(self.errors, row, column, problem)

    def add_warning(self, row: int, column: str, problem: str) -> None:
        self.warning_count += 1
        self._add(self.warnings, row, column, problem)

    def summary(self) -> dict[str, Any]:
        """The counts and the first problems, as shown to the user."""
        return {
            "file": Path(self.path).name,
            "format": self.format,
            "source": self.source,
            "rows_read": self.rows_read,
            "valid_rows": self.valid_rows,
            "rows_skipped": self.rows_skipped,
            "duplicates_merged": self.duplicates,
            "contacts": len(self.contacts),
            "error_count": self.error_count,
            "warning_count": self.warning_count,
            "errors": self.errors[:SHOWN],
            "warnings": self.warnings[:SHOWN],
            "unknown_columns": self.unknown_columns,
        }


# --- reading a file ---------------------------------------------------------------------------


def _header_key(text: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(text or "").strip().casefold()).strip("_")


def _find_header(
    rows: Iterable[tuple[int, list[Any]]],
) -> tuple[int, dict[int, str], list[str]] | None:
    """The first row that names at least two known columns (one of them about the person).
    Returns `(row number, {cell index: column}, unknown column names)`."""
    for number, cells in rows:
        mapped: dict[int, str] = {}
        taken: set[str] = set()
        unknown: list[str] = []
        for index, cell in enumerate(cells):
            column = _ALIASES.get(_header_key(cell))
            if column is not None and column not in taken:
                mapped[index] = column
                taken.add(column)
            elif str(cell or "").strip():
                unknown.append(str(cell).strip()[:60])
        if len(mapped) >= 2 and taken & set(_IDENTITY):
            return number, mapped, unknown
        if number > HEADER_SCAN_ROWS:
            break
    return None


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _csv_rows(text: str, delimiter: str) -> Iterator[tuple[int, list[Any]]]:
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    for cells in reader:
        yield reader.line_num, cells


def _csv_source(path: Path) -> tuple[Iterator[tuple[int, list[Any]]], Callable[[], None]]:
    text = _decode(path.read_bytes()).replace("\x00", "")
    best: tuple[int, str] | None = None
    for delimiter in (",", ";", "\t"):
        try:
            found = _find_header(_csv_rows(text, delimiter))
        except csv.Error:
            continue
        if found and (best is None or len(found[1]) > best[0]):
            best = (len(found[1]), delimiter)
    delimiter = best[1] if best else ","
    return _csv_rows(text, delimiter), lambda: None


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _xlsx_source(path: Path) -> tuple[Iterator[tuple[int, list[Any]]], Callable[[], None]]:
    import openpyxl  # imported late: only Excel files need it

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = {name.casefold(): name for name in workbook.sheetnames}
    sheet = workbook[sheets.get(SHEET_NAME.casefold(), workbook.sheetnames[0])]
    sheet.reset_dimensions()  # do not trust the size a tool wrote into the file

    def rows() -> Iterator[tuple[int, list[Any]]]:
        blank_run = 0
        for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = [_cell(v) for v in values]
            if any(str(c).strip() for c in cells):
                blank_run = 0
            else:
                blank_run += 1
                if blank_run >= MAX_BLANK_RUN or number >= MAX_SCANNED_ROWS:
                    return
            yield number, cells

    return rows(), workbook.close


# --- cleaning ---------------------------------------------------------------------------------


def _text(value: Any, limit: int) -> str:
    return normalize_whitespace(_BAD_CELL_RE.sub("", str(value or "")))[:limit]


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    try:  # 2026-09-30T10:00:00 and similar
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError:
        return None


def _linkedin(value: str) -> str:
    text = value.strip()
    if text and "://" not in text and "linkedin.com/" in text.lower():
        text = "https://" + text.lstrip("/")
    return canonical_url(text)


def _split_name(full: str) -> tuple[str, str]:
    parts = full.split(" ", 1)
    return parts[0], parts[1] if len(parts) > 1 else ""


def _clean_row(
    number: int, raw: Mapping[str, Any], preview: ImportPreview, today: date
) -> Contact | None:
    """One row as a `Contact`, or None when it cannot be used (the reason is in the preview)."""
    first = _text(raw.get("first_name"), _LIMITS["first_name"])
    last = _text(raw.get("last_name"), _LIMITS["last_name"])
    if not first and not last:
        first, last = _split_name(_text(raw.get("full_name"), _LIMITS["full_name"]))
    contact = Contact(
        first_name=first, last_name=last, rows=[number],
        company=_text(raw.get("company"), _LIMITS["company"]),
        title=_text(raw.get("title"), _LIMITS["title"]),
        country=_text(raw.get("country"), _LIMITS["country"]),
        notes=_text(raw.get("notes"), _LIMITS["notes"]),
    )  # fmt: skip

    email = str(raw.get("email") or "").strip()
    if email.casefold().startswith("mailto:"):
        email = email[7:]
    if email:
        contact.email = suppression.normalize(email, "email")
        if not contact.email:
            preview.add_error(
                number, "email", f"'{_text(email, 60)}' is not a valid email address."
            )
    link = _text(raw.get("linkedin_url"), 300)
    if link:
        contact.linkedin_url = _linkedin(link)
        if not contact.linkedin_url:
            preview.add_error(
                number, "linkedin_url",
                f"'{link[:60]}' is not a LinkedIn profile link "
                "(it must look like https://www.linkedin.com/in/name).",
            )  # fmt: skip
    domain = _text(raw.get("company_domain"), 200)
    if domain:
        contact.company_domain = suppression.normalize(domain, "domain")
        if not contact.company_domain:
            preview.add_error(number, "company_domain", f"'{domain[:60]}' is not a valid domain.")

    channel = _header_key(raw.get("channel"))
    channel = _CHANNEL_ALIASES.get(channel, channel)
    if channel and channel not in CHANNELS:
        preview.add_error(number, "channel", f"'{channel[:30]}' is not email, linkedin or both.")
        channel = ""
    contact.channel = channel

    relationship = _header_key(raw.get("relationship"))
    relationship = _RELATIONSHIP_ALIASES.get(relationship, relationship) or "contacted"
    if relationship not in RELATIONSHIPS:
        preview.add_error(
            number, "relationship",
            f"'{relationship[:30]}' is not one of: {', '.join(RELATIONSHIPS)}. "
            "The row was skipped so that nobody is treated more gently than you meant.",
        )  # fmt: skip
        return None
    contact.relationship = relationship

    parsed: date | None = None
    given = raw.get("last_contacted")
    if given is not None and str(given).strip():
        parsed = _parse_date(given)
        if parsed is None or parsed.year < _MIN_YEAR:
            preview.add_error(
                number, "last_contacted",
                f"'{_text(given, 30)}' is not a date (use 2026-09-30, 30/09/2026 or 30-09-2026). "
                "The row was skipped.",
            )  # fmt: skip
            return None
        if parsed > today:
            preview.add_warning(
                number, "last_contacted", "The date is in the future; today is used."
            )
            parsed = today
    contact.last_contacted = parsed.isoformat() if parsed else ""

    keys = contact.keys()
    if not any(keys.values()):
        preview.add_error(
            number, "row",
            "Needs an email, a LinkedIn link, or a first/last name together with a company.",
        )
        return None
    if not contact.email and not contact.linkedin_url:
        preview.add_warning(
            number, "row", "No email and no LinkedIn link: can only be matched by name and company."
        )
    return contact


def _stronger(a: str, b: str) -> str:
    return a if _STRONGEST.index(a) >= _STRONGEST.index(b) else b


def _merge(into: Contact, other: Contact) -> None:
    for name in (
        "first_name", "last_name", "email", "linkedin_url", "company", "company_domain",
        "title", "country",
    ):  # fmt: skip
        if not getattr(into, name):
            setattr(into, name, getattr(other, name))
    into.last_contacted = max(into.last_contacted, other.last_contacted)
    into.relationship = _stronger(into.relationship, other.relationship)
    if into.channel != other.channel and other.channel:
        into.channel = other.channel if not into.channel else "both"
    if other.notes and other.notes not in into.notes:
        into.notes = _text(f"{into.notes} | {other.notes}" if into.notes else other.notes, 500)
    into.rows.extend(other.rows)


class _Merger:
    """Merges rows that share a LinkedIn link, an email, or a name and company."""

    def __init__(self) -> None:
        self.contacts: list[Contact] = []
        self._index: dict[tuple[str, str], Contact] = {}

    def add(self, contact: Contact) -> bool:
        """Add a row; True when it was merged into an earlier one."""
        keys = [(k, v) for k, v in contact.keys().items() if v]
        found = next((self._index[key] for key in keys if key in self._index), None)
        target = found or contact
        if found is not None:
            _merge(found, contact)
        else:
            self.contacts.append(contact)
        for kind, value in target.keys().items():
            if value:
                self._index.setdefault((kind, value), target)
        return found is not None


def _linkedin_notes(raw: Mapping[str, Any]) -> str:
    connected = _parse_date(raw.get("connected_on"))
    return f"LinkedIn connection since {connected.isoformat()}." if connected else ""


def _read_rows(
    preview: ImportPreview, rows: Iterator[tuple[int, list[Any]]], today: date
) -> None:
    header = _find_header(rows)
    # `_find_header` consumed rows from the iterator up to the header; the rest follow.
    if header is None:
        preview.error = (
            "No header row found. The first row must name the columns, for example "
            f"{', '.join(COLUMNS[:6])}, ... Download the template (get_contacts_template)."
        )
        return
    _, mapped, preview.unknown_columns = header
    if "connected_on" in mapped.values() and {"first_name", "linkedin_url"} <= set(mapped.values()):
        preview.source = "linkedin_connections"
    merger = _Merger()
    for number, cells in rows:
        raw = {column: cells[i] for i, column in mapped.items() if i < len(cells)}
        if not any(str(v).strip() for v in raw.values()):
            continue  # a blank row (or one holding only unknown columns)
        preview.rows_read += 1
        if preview.rows_read > MAX_ROWS:
            preview.error = (
                f"The file has more than {MAX_ROWS:,} rows. Split it into smaller files "
                f"(up to {MAX_ROWS:,} rows each) and import them one by one."
            )
            return
        if preview.source == "linkedin_connections":
            raw = {**raw, "notes": _linkedin_notes(raw)}
        contact = _clean_row(number, raw, preview, today)
        if contact is None:
            preview.rows_skipped += 1
            continue
        preview.valid_rows += 1
        if merger.add(contact):
            preview.duplicates += 1
    preview.contacts = merger.contacts


def read_contacts(path: str | Path) -> ImportPreview:
    """Read a `.csv` or `.xlsx` file into an `ImportPreview`. Never raises: a file that cannot
    be used has `error` set, bad rows are listed in `errors` / `warnings`."""
    file = Path(str(path)).expanduser()
    preview = ImportPreview(path=str(file), format=file.suffix.lower().lstrip("."))
    try:
        if not file.is_file():
            preview.error = f"There is no file at {file}."
            return preview
        if file.suffix.lower() not in (".csv", ".xlsx"):
            preview.error = (
                f"'{file.name}' is a {file.suffix or 'file without an extension'} file. "
                "Save it as .csv or .xlsx (in Excel: File, Save As) and try again."
            )
            return preview
        if file.stat().st_size > MAX_BYTES:
            preview.error = (
                f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB. "
                "Split it into smaller files."
            )
            return preview
        opener = _xlsx_source if preview.format == "xlsx" else _csv_source
        rows, close = opener(file)
        try:
            _read_rows(preview, rows, datetime.now(UTC).date())
        finally:
            close()
    except Exception as exc:  # a messy real-world file must never crash the tool
        kind = "Excel" if preview.format == "xlsx" else "CSV"
        preview.error = (
            f"The file could not be read as a {kind} file ({type(exc).__name__}). "
            "Open it, save a copy as .csv (or .xlsx) and try again."
        )
    if preview.error:
        preview.contacts = []
    return preview


# --- template ---------------------------------------------------------------------------------


def template_dir() -> Path:
    """The workspace folder `templates/` (created) where `template_info` puts the files."""
    folder = paths.home() / "templates"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def _csv_template_bytes() -> bytes:
    return resources.files("leadgen").joinpath("templates", TEMPLATE_CSV).read_bytes()


def export_template(path: str | Path, fmt: str = "csv") -> Path:
    """Write the empty template as CSV (UTF-8 with BOM) or Excel (two sheets) and return its path.

    `path` is a file, or an existing folder (the standard file name is used). Raises `ValueError`
    for another `fmt`."""
    fmt = fmt.lower()
    if fmt not in ("csv", "xlsx"):
        raise ValueError("fmt must be csv or xlsx")
    target = Path(path).expanduser()
    if target.is_dir():
        target = target / (TEMPLATE_CSV if fmt == "csv" else TEMPLATE_XLSX)
    target.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "csv":
        target.write_bytes(_csv_template_bytes())
        return target
    import openpyxl
    from openpyxl.styles import Alignment, Font

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(list(COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, name in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = max(
            16, len(name) + 4
        )
    sheet.freeze_panes = "A2"
    guide = workbook.create_sheet(GUIDE_SHEET_NAME)
    guide.append(["Column", "Needed?", "What to put there", "Example"])
    for cell in guide[1]:
        cell.font = Font(bold=True)
    for name in COLUMNS:
        needed, meaning, example = COLUMN_GUIDE[name]
        guide.append([name, needed, meaning, example])
    guide.append([])
    guide.append([_NEEDS])
    for letter, width in (("A", 18), ("B", 16), ("C", 90), ("D", 36)):
        guide.column_dimensions[letter].width = width
    for row in guide.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(target)
    return target


def template_info() -> dict[str, Any]:
    """Write both templates into the workspace and describe the columns."""
    folder = template_dir()
    csv_path = export_template(folder / TEMPLATE_CSV, "csv")
    xlsx_path = export_template(folder / TEMPLATE_XLSX, "xlsx")
    return {
        "ok": True,
        "csv_path": str(csv_path),
        "xlsx_path": str(xlsx_path),
        "columns": [
            {"name": n, "needed": COLUMN_GUIDE[n][0], "meaning": COLUMN_GUIDE[n][1],
             "example": COLUMN_GUIDE[n][2]}
            for n in COLUMNS
        ],  # fmt: skip
        "rules": _NEEDS,
        "relationship_values": list(RELATIONSHIPS),
        "date_formats": ["2026-09-30", "30/09/2026", "30-09-2026"],
        "linkedin_tip": (
            "LinkedIn's own export also works as it is: Settings, Data privacy, Get a copy of "
            "your data, Connections. Use the Connections.csv file from the download."
        ),
        "next_step": (
            "Give the user the file path. When they have filled it in, call import_contacts "
            "with dry_run=true and show them the preview before applying."
        ),
    }


# --- importing --------------------------------------------------------------------------------


def _check_mode(mode: str) -> str:
    if mode not in MODES:
        raise ContactsImportError(
            "invalid_mode", f"mode must be one of: {', '.join(MODES)}.",
            "exclude holds these people back for the cooldown; signal only flags them as known; "
            "both does both.",
        )  # fmt: skip
    return mode


def _store_known(
    conn: sqlite3.Connection, project_id: int, contacts: Sequence[Contact], source: str
) -> dict[str, int]:
    """Insert or update `known_contacts`. A contact already stored (same LinkedIn, else email,
    else name and company) is updated, so importing a file again adds nothing."""
    stamp = db.now()
    new = updated = unchanged = 0
    for contact in contacts:
        keys = contact.keys()
        company_key = name_key(contact.company)
        existing = None
        for kind in ("key_linkedin", "key_email", "key_name_company"):
            if keys[kind]:
                existing = conn.execute(
                    f"SELECT * FROM known_contacts WHERE project_id = ? AND {kind} = ? "
                    "ORDER BY id LIMIT 1",
                    (project_id, keys[kind]),
                ).fetchone()
                if existing is not None:
                    break
        values = {
            "first_name": contact.first_name, "last_name": contact.last_name,
            "company": contact.company, "company_domain": contact.company_domain,
            "title": contact.title, "country": contact.country, "channel": contact.channel,
            "notes": contact.notes,
        }  # fmt: skip
        if existing is None:
            db.insert(
                conn, "known_contacts", project_id=project_id, key_company=company_key,
                relationship=contact.relationship, last_contacted=contact.last_contacted,
                source=source[:200], imported_at=stamp, updated_at=stamp, **keys, **values,
            )  # fmt: skip
            new += 1
            continue
        change: dict[str, Any] = {}
        for column, value in {**keys, "key_company": company_key, **values}.items():
            if value and existing[column] != value:
                change[column] = value
        relationship = _stronger(
            existing["relationship"] if existing["relationship"] in RELATIONSHIPS else "contacted",
            contact.relationship,
        )
        if relationship != existing["relationship"]:
            change["relationship"] = relationship
        latest = max(existing["last_contacted"], contact.last_contacted)
        if latest != existing["last_contacted"]:
            change["last_contacted"] = latest
        if change:
            db.update(conn, "known_contacts", existing["id"], updated_at=stamp, **change)
            updated += 1
        else:
            unchanged += 1
    return {"new": new, "updated": updated, "unchanged": unchanged}


def _suppress(
    conn: sqlite3.Connection, contacts: Sequence[Contact], source: str
) -> tuple[dict[str, int], int]:
    """Put the do-not-contact contacts' emails and LinkedIn links on the global list.
    Returns `({"added", "already_listed"}, name_only)`."""
    emails = [c.email for c in contacts if c.email]
    links = [c.linkedin_url for c in contacts if c.linkedin_url]
    name_only = sum(1 for c in contacts if not c.email and not c.linkedin_url)
    added = listed = 0
    reason = f"do not contact ({source})"
    for kind, values in (("email", emails), ("linkedin", links)):
        for start in range(0, len(values), suppression.MAX_VALUES):
            result = suppression.add_in(
                conn, values[start : start + suppression.MAX_VALUES], kind, reason
            )
            added += len(result["added"])
            listed += len(result["already_listed"])
    return {"added": added, "already_listed": listed}, name_only


def _apply(
    conn: sqlite3.Connection,
    project_id: int,
    preview: ImportPreview,
    mode: str,
    source: str,
    cooldown_days: int,
) -> dict[str, Any]:
    today = datetime.now(UTC).date().isoformat()
    contacts = preview.contacts
    blocked = [c for c in contacts if c.relationship == "do_not_contact"]
    out: dict[str, Any] = {"do_not_contact": len(blocked)}
    if blocked:
        listed, name_only = _suppress(conn, blocked, source)
        out["suppressed"] = listed["added"]
        out["already_suppressed"] = listed["already_listed"]
        if name_only:
            out["do_not_contact_name_only"] = name_only
    if mode in ("exclude", "both"):
        rows = [{**c.as_person(), "at": c.last_contacted or today} for c in contacts]
        logged = projects.log_contacts_in(
            conn, project_id, rows, "imported", None, source, None
        )
        out["excluded"] = logged["logged"]
        out["already_logged"] = logged["skipped"]
        cutoff = (datetime.now(UTC) - timedelta(days=cooldown_days)).date().isoformat()
        out["outside_cooldown"] = sum(
            1 for c in contacts if c.last_contacted and c.last_contacted < cutoff
        )
    if mode in ("signal", "both"):
        out["signalled"] = _store_known(
            conn, project_id, [c for c in contacts if c.relationship != "do_not_contact"], source
        )
    return out


def import_contacts(
    project_id: int,
    path: str | Path,
    mode: str = "exclude",
    dry_run: bool = True,
    file_root: Path | None = None,
) -> dict[str, Any]:
    """Read a contacts file and use it in `mode` (`exclude`, `signal` or `both`).

    With `dry_run` (the default) everything is worked out inside a transaction that is rolled
    back: the counts are exact and nothing is written. `file_root` (remote mode) limits the file
    to that folder. Raises `ContactsImportError` (`invalid_mode`, `unreadable_file`,
    `file_outside_workspace`) and `projects.ProjectError` (`not_found`).
    """
    _check_mode(mode)
    file = Path(str(path)).expanduser()
    if file_root is not None:
        try:
            inside = file.resolve().is_relative_to(file_root.resolve())
        except OSError:
            inside = False
        if not inside:
            raise ContactsImportError(
                "file_outside_workspace",
                "In this connection only files inside the Lead Generator folder can be read.",
                f"Ask the user to copy the file into {file_root} (the templates folder is fine).",
            )
    project = projects.get_project(project_id)
    preview = read_contacts(file)
    if preview.error:
        raise ContactsImportError(
            "unreadable_file", preview.error,
            "Tell the user what is wrong with the file; get_contacts_template gives a clean "
            "template.",
        )  # fmt: skip
    source = f"import: {file.name}"
    conn = db.connect()
    try:
        applied = _apply(conn, project_id, preview, mode, source, project["cooldown_days"])
        if dry_run:
            conn.rollback()
        else:
            conn.commit()
    except BaseException:
        conn.rollback()
        raise
    finally:
        conn.close()
    return _reply(project, preview, mode, dry_run, applied)


def _reply(
    project: Mapping[str, Any],
    preview: ImportPreview,
    mode: str,
    dry_run: bool,
    applied: dict[str, Any],
) -> dict[str, Any]:
    summary = preview.summary()
    out: dict[str, Any] = {
        "ok": True, "dry_run": dry_run, "mode": mode, "project_id": project["id"],
    }  # fmt: skip
    out.update(summary)
    out["do_not_contact"] = applied["do_not_contact"]
    if "excluded" in applied:
        out["would_exclude" if dry_run else "excluded"] = applied["excluded"]
        out["already_logged"] = applied["already_logged"]
        out["outside_cooldown"] = applied["outside_cooldown"]
    if "signalled" in applied:
        known = applied["signalled"]
        out["would_signal" if dry_run else "signalled"] = known["new"] + known["updated"]
        out["signal_new"], out["signal_updated"] = known["new"], known["updated"]
    for key in ("suppressed", "already_suppressed", "do_not_contact_name_only"):
        if key in applied:
            out[key] = applied[key]
    out["message"] = _message(project, summary, mode, dry_run, applied)
    out["next_step"] = (
        "Show the user these counts and the first errors. If they are happy, call import_contacts "
        "again with dry_run=false."
        if dry_run
        else "Done. Tell the user what was imported; new campaigns of this project now respect it."
    )
    return out


def _message(
    project: Mapping[str, Any],
    summary: Mapping[str, Any],
    mode: str,
    dry_run: bool,
    applied: Mapping[str, Any],
) -> str:
    head = "Preview only, nothing was written. " if dry_run else ""
    parts = [
        f"{summary['rows_read']} rows read, {summary['contacts']} different people"
        f" ({summary['duplicates_merged']} repeated rows merged, {summary['rows_skipped']} rows "
        f"skipped for errors, {summary['error_count']} problems, {summary['warning_count']} "
        "warnings)."
    ]
    verb = "would be" if dry_run else "were"
    if "excluded" in applied:
        parts.append(
            f"{applied['excluded']} {verb} held back in project '{project['name']}' for its "
            f"{project['cooldown_days']}-day cooldown ({applied['already_logged']} already there)."
        )
        if applied["outside_cooldown"]:
            parts.append(
                f"{applied['outside_cooldown']} were last contacted longer ago than the cooldown, "
                "so they do not hold anyone back now (mark people who must never be contacted as "
                "do_not_contact)."
            )
    if "signalled" in applied:
        known = applied["signalled"]
        parts.append(
            f"{known['new'] + known['updated']} {verb} saved as known contacts "
            "(their companies count as warm)."
        )
    if applied["do_not_contact"]:
        parts.append(
            f"{applied['do_not_contact']} marked do not contact: their emails and LinkedIn links "
            f"{verb} added to the do-not-contact list."
        )
    return head + " ".join(parts)


# --- using the stored contacts ----------------------------------------------------------------

_CHUNK = 300


def _chunks(values: Sequence[str]) -> Iterator[list[str]]:
    for start in range(0, len(values), _CHUNK):
        yield list(values[start : start + _CHUNK])


def known_for_people(
    project_id: int | None, person_ids: Sequence[int]
) -> dict[int, dict[str, str]]:
    """`{person_id: {"relationship", "last_contacted"}}` for the people who are stored known
    contacts of the project (matched by LinkedIn, then email, then name and company)."""
    ids = list(dict.fromkeys(int(i) for i in person_ids))
    if project_id is None or not ids:
        return {}
    with db.session() as conn:
        if not conn.execute(
            "SELECT 1 FROM known_contacts WHERE project_id = ? LIMIT 1", (project_id,)
        ).fetchone():
            return {}
        people = db.query(
            conn,
            "SELECT p.id, p.linkedin_url, p.first_name, p.last_name, p.raw_name, "
            "c.name AS company FROM people p JOIN companies c ON c.id = p.company_id "
            f"WHERE p.id IN ({', '.join('?' for _ in ids)})",
            ids,
        )
        addresses: dict[int, str] = {}
        for row in db.query(
            conn,
            "SELECT person_id, email FROM emails WHERE person_id IN "
            f"({', '.join('?' for _ in ids)}) AND status <> 'invalid' ORDER BY id DESC",
            ids,
        ):
            addresses[row["person_id"]] = str(row["email"]).lower()
        wanted = {p["id"]: projects.make_keys({**p, "email": addresses.get(p["id"], "")})
                  for p in people}  # fmt: skip
        found: dict[tuple[str, str], dict[str, Any]] = {}
        for kind in ("key_linkedin", "key_email", "key_name_company"):
            values = sorted({k[kind] for k in wanted.values() if k[kind]})
            for chunk in _chunks(values):
                for row in db.query(
                    conn,
                    "SELECT id, relationship, last_contacted, "
                    f"{kind} AS key FROM known_contacts WHERE project_id = ? AND {kind} IN "
                    f"({', '.join('?' for _ in chunk)}) ORDER BY id",
                    [project_id, *chunk],
                ):
                    found.setdefault((kind, row["key"]), row)
    out: dict[int, dict[str, str]] = {}
    for person_id, keys in wanted.items():
        for kind in ("key_linkedin", "key_email", "key_name_company"):
            row = found.get((kind, keys[kind])) if keys[kind] else None
            if row is not None:
                out[person_id] = {
                    "relationship": row["relationship"],
                    "last_contacted": row["last_contacted"],
                }
                break
    return out


def warm_counts(
    project_id: int | None, companies: Sequence[Mapping[str, Any]]
) -> dict[int, int]:
    """`{company_id: number of known contacts}` for companies with at least one. A company
    matches by its name, or by its domain when the import gave one."""
    if project_id is None or not companies:
        return {}
    out: dict[int, int] = {}
    with db.session() as conn:
        if not conn.execute(
            "SELECT 1 FROM known_contacts WHERE project_id = ? LIMIT 1", (project_id,)
        ).fetchone():
            return {}
        for company in companies:
            key = name_key(str(company.get("name") or ""))
            domain = str(company.get("domain") or "").lower()
            count = conn.execute(
                "SELECT COUNT(*) FROM known_contacts WHERE project_id = ? AND "
                "((? <> '' AND key_company = ?) OR (? <> '' AND company_domain = ?))",
                (project_id, key, key, domain, domain),
            ).fetchone()[0]
            if count:
                out[int(company["id"])] = int(count)
    return out
